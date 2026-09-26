from __future__ import annotations

import argparse
import html
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook

SOURCE = "https://www.gov.br/anvisa/pt-br/assuntos/medicamentos/cmed/precos"
RATE_PATTERN = re.compile(r"^(PF|PMVG)\s+(\d+(?:,\d+)?)\s*%$")


def price(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2) if value >= 0 else None
    cleaned = str(value).strip().replace("*", "").replace(".", "").replace(",", ".")
    try:
        result = float(cleaned)
        return round(result, 2) if result >= 0 else None
    except ValueError:
        return None


def rate_columns(header):
    columns = {}
    for index, value in enumerate(header):
        match = RATE_PATTERN.match(str(value or "").strip())
        if match:
            kind, rate = match.groups()
            columns.setdefault(rate, {})[kind.lower()] = index
    return {rate: pair for rate, pair in columns.items() if set(pair) == {"pf", "pmvg"}}


def rows_from_xlsx(path):
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    if sheet is None:
        raise ValueError("Planilha vazia")
    iterator = sheet.iter_rows(values_only=True)
    for header_row, header in enumerate(iterator, 1):
        if header[0] == "SUBSTÂNCIA" and header[3] == "CÓDIGO GGREM":
            break
    else:
        raise ValueError("Cabeçalho CMED não encontrado")
    rates = rate_columns(header)
    if "0" not in rates:
        raise ValueError("Colunas PF/PMVG 0% não encontradas")
    found, rejected = [], 0
    for row in iterator:
        ggrem = str(row[3] or "").strip()
        if not ggrem:
            continue
        values = {}
        for rate, pair in rates.items():
            raw_pf, raw_pmvg = row[pair["pf"]], row[pair["pmvg"]]
            values[rate] = [price(raw_pf), price(raw_pmvg),
                            "*" in str(raw_pf or "") or "*" in str(raw_pmvg or "")]
        if values["0"][0] is None:
            rejected += 1
            continue
        found.append({"g": ggrem, "s": str(row[0] or "").strip(),
                      "l": str(row[2] or "").strip(), "n": str(row[8] or "").strip(),
                      "a": str(row[9] or "").strip(), "rates": values,
                      "cap": str(row[66] or "").strip(), "hospital": str(row[65] or "").strip()})
    book.close()
    if not found:
        raise ValueError("Nenhum registro válido")
    return found, rejected, header_row, list(rates)


def build(source, output):
    records, rejected, header_row, rates = rows_from_xlsx(source)
    if len({r["g"] for r in records}) != len(records):
        raise ValueError("GGREM duplicado: revisão da origem necessária")
    report = {"source_file": source.name, "source_url": SOURCE,
              "generated_utc": datetime.now(timezone.utc).isoformat(),
              "records": len(records), "unique_ggrem": len(records),
              "rejected_without_pf_zero": rejected, "header_row": header_row,
              "rates": rates, "cap_marked": sum(r["cap"].casefold() == "sim" for r in records),
              "hospital_restricted": sum(r["hospital"].casefold() == "sim" for r in records)}
    strings, indexes, packed_rows = [], {}, []
    def intern(value):
        if value not in indexes:
            indexes[value] = len(strings)
            strings.append(value)
        return indexes[value]
    for row in records:
        packed_rows.append([row["g"], intern(row["s"]), intern(row["l"]), intern(row["n"]),
                            intern(row["a"]), row["cap"], row["hospital"],
                            [row["rates"][rate] for rate in rates]])
    packed = {"strings": strings, "rows": packed_rows}
    payload = json.dumps(packed, ensure_ascii=False, separators=(",", ":")).replace("<", "\\u003c")
    page = Path(__file__).with_name("template.html").read_text(encoding="utf-8")
    page = page.replace("__DATA__", payload).replace("__RATES__", json.dumps(rates))
    page = page.replace("__DATE__", html.escape(source.stem.replace("cmed_", "")))
    output.write_text(page, encoding="utf-8")
    output.with_suffix(".quality.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Gera console local de conferência CMED")
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, default=Path("dashboard.html"))
    args = parser.parse_args()
    print(json.dumps(build(args.source, args.output), indent=2, ensure_ascii=False))
